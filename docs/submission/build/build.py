# -*- coding: utf-8 -*-
import re, os, sys
sys.path.insert(0, os.path.dirname(__file__))
from content import *
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(BASE, 'fig')
KAI, TNR = '標楷體', 'Times New Roman'
BODY, SMALL, LINE = 14, 12, 18

# ---- citation renumbering by first appearance ----
order = []
def cite(text):
    def rep(m):
        k = m.group(1)
        if k not in order: order.append(k)
        return '[%d]' % (order.index(k)+1)
    t = re.sub(r'\{(R\d+)\}', rep, text)
    t = re.sub(r'\](\s*)\[', '], [', t)   # [1][2] -> [1], [2]
    return t.replace('], , [', '], [')

def set_font(run, size=BODY, bold=False, italic=False):
    run.font.name = TNR; run.font.size = Pt(size); run.bold = bold; run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii','w:hAnsi','w:cs'): rf.set(qn(a), TNR)
    rf.set(qn('w:eastAsia'), KAI)

def fmt_par(p, size=BODY, align=None, first=None, left=None, hanging=None, before=0, after=0, keep=False):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY; pf.line_spacing = Pt(LINE)
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if align is not None: p.alignment = align
    if left is not None: pf.left_indent = Cm(left)
    if first is not None: pf.first_line_indent = Cm(first)
    if hanging is not None: pf.first_line_indent = Cm(-hanging)
    if keep: pf.keep_with_next = True

def add_text(p, text, size=BODY, bold=False):
    r = p.add_run(cite(text)); set_font(r, size, bold); return r

def cell_par(cell, first=False):
    if first and cell.paragraphs and not cell.paragraphs[0].text and len(cell.paragraphs)==1 and not cell.paragraphs[0].runs:
        return cell.paragraphs[0]
    return cell.add_paragraph()

def set_cell_width(cell, cm):
    tcPr = cell._tc.get_or_add_tcPr()
    w = tcPr.find(qn('w:tcW'))
    if w is None: w = OxmlElement('w:tcW'); tcPr.append(w)
    w.set(qn('w:w'), str(int(cm*567))); w.set(qn('w:type'), 'dxa')

def set_cell_shading(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexcolor); tcPr.append(sh)

def table_borders(tbl, sz=8):
    tblPr = tbl._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x = OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),str(sz)); x.set(qn('w:space'),'0'); x.set(qn('w:color'),'000000'); b.append(x)
    tblPr.append(b)

def set_col_widths(tbl, widths):
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    grid = tbl._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(int(widths[i]*567)))
    for row in tbl.rows:
        for i, c in enumerate(row.cells): set_cell_width(c, widths[i])

def nob(cell, sides):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for s in sides:
        e = OxmlElement('w:'+s); e.set(qn('w:val'),'nil'); b.append(e)
    tcPr.append(b)

def cant_split(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement('w:cantSplit'); trPr.append(e)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement('w:tblHeader'); trPr.append(e)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
sec.footer_distance = Cm(1.0)
st = doc.styles['Normal']; st.font.name = TNR; st.font.size = Pt(BODY)
rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
for a in ('w:ascii','w:hAnsi','w:cs'): rf.set(qn(a), TNR)
rf.set(qn('w:eastAsia'), KAI)
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; st.paragraph_format.line_spacing = Pt(LINE)
st.paragraph_format.space_after = Pt(0)
# lang zh-TW
lang = OxmlElement('w:lang'); lang.set(qn('w:val'),'en-US'); lang.set(qn('w:eastAsia'),'zh-TW'); rpr.append(lang)

# footer page number
fp = sec.footer.paragraphs[0]; fmt_par(fp, align=WD_ALIGN_PARAGRAPH.CENTER)
r = fp.add_run(); set_font(r, 12)
for tag, txt in (('begin',None),(None,'PAGE'),('end',None)):
    if tag:
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), tag); r._r.append(e)
    else:
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text = txt; r._r.append(it)

# core props
doc.core_properties.title = TITLE; doc.core_properties.author = ''; doc.core_properties.comments = ''

p = doc.add_paragraph(); fmt_par(p); add_text(p, '【附件1-2】', BODY, True)
p = doc.add_paragraph(); fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6); add_text(p, '提案資料表（研究實作組）', 16, True)

LW, CW = 3.0, 14.0
T = doc.add_table(rows=0, cols=2); table_borders(T, 8); T.alignment = WD_TABLE_ALIGNMENT.CENTER

def row(label, short=False):
    rw = T.add_row(); c0, c1 = rw.cells
    p0 = c0.paragraphs[0]; fmt_par(p0, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, part in enumerate(label.split('\n')):
        if i: p0.add_run().add_break()
        add_text(p0, part, BODY, True)
    c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(c0, 'F2F2F2')
    if short: cant_split(rw)
    return c1

def section(label):
    rw = T.add_row(); a, b = rw.cells; m = a.merge(b)
    p0 = m.paragraphs[0]; fmt_par(p0, align=WD_ALIGN_PARAGRAPH.CENTER); add_text(p0, label.replace('\n',''), BODY, True)
    set_cell_shading(m, 'F2F2F2'); cant_split(rw)
    trPr = rw._tr.get_or_add_trPr(); 
    rw2 = T.add_row(); a, b = rw2.cells; m2 = a.merge(b)
    return m2

def blocks(cell, items, size=BODY):
    first = True
    for kind, val in items:
        if kind == 'table': inner_table(cell, TABLES[val]); first=False; continue
        if kind == 'fig': figure(cell, *FIGS[val]); first=False; continue
        p = cell_par(cell, first); first = False
        if kind == 'h1': fmt_par(p, before=4); add_text(p, val, size, True)
        elif kind == 'h2': fmt_par(p, left=0.0); add_text(p, val, size, True)
        elif kind == 'p': fmt_par(p, first=0.99, align=WD_ALIGN_PARAGRAPH.JUSTIFY); add_text(p, val, size)
        elif kind == 'li':
            fmt_par(p, left=0.6, hanging=0.6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            m = re.match(r'^(\d+\.)(.*)$', val)
            add_text(p, m.group(1) if m else '', size); add_text(p, m.group(2) if m else val, size)

def inner_table(cell, spec):
    ncol = len(spec['header'])
    sp = cell.add_paragraph(); fmt_par(sp)
    t = cell.add_table(rows=2+len(spec['rows']), cols=ncol); table_borders(t, 4)
    total = LW + CW - 0.6
    if spec.get('widths'): w = spec['widths']
    elif ncol == 5: w = [3.0, 2.9, 2.4, 2.6, 2.7]
    elif spec['header'][0] == '指標': w = [5.6, 2.2, 2.6, 3.2]
    else: w = [3.6, 2.8, 3.8, 3.8]
    s = sum(w); w = [x*total/s for x in w]
    set_col_widths(t, w)
    cap = t.rows[0].cells[0].merge(t.rows[0].cells[ncol-1])
    pp = cap.paragraphs[0]; fmt_par(pp, align=WD_ALIGN_PARAGRAPH.CENTER); add_text(pp, spec['caption'], SMALL, True)
    nob(cap, ('top','left','right')); cant_split(t.rows[0])
    for i, h in enumerate(spec['header']):
        c = t.rows[1].cells[i]; pp = c.paragraphs[0]; fmt_par(pp, align=WD_ALIGN_PARAGRAPH.CENTER); add_text(pp, h, SMALL, True); set_cell_shading(c, 'E7E6E6')
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    repeat_header(t.rows[0]); repeat_header(t.rows[1]); cant_split(t.rows[1])
    for ri, rr in enumerate(spec['rows']):
        cant_split(t.rows[ri+2])
        for i, v in enumerate(rr):
            c = t.rows[ri+2].cells[i]; pp = c.paragraphs[0]
            fmt_par(pp, align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            bold = ('可再配置' in rr[0]) or ('合計' in rr[0])
            add_text(pp, v, SMALL, bold)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for rw in t.rows[:-1]:
        for cc in rw.cells:
            for pp in cc.paragraphs: pp.paragraph_format.keep_with_next = True
    p = cell.add_paragraph(); fmt_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6); add_text(p, spec['note'], SMALL)

def figure(cell, fname, width, caption):
    p = cell.add_paragraph(); fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, keep=True)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE  # image paragraph must not be clipped by exact spacing
    p.add_run().add_picture(os.path.join(FIG, fname), width=Cm(width))
    p = cell.add_paragraph(); fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=6); add_text(p, caption, SMALL, True)

def simple(cell, text, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]; fmt_par(p, align=align); add_text(p, text); return p

c = row('報名編號', True); simple(c, '（本欄位由主辦單位填寫）')
c = row('團隊名稱', True); simple(c, '＿＿＿＿＿＿＿＿＿＿')
c = row('指導教師', True); simple(c, '姓名：＿＿＿＿＿；學校/系所(單位)：＿＿＿＿＿＿＿＿')
c = row('團隊代表人', True); simple(c, '姓名：＿＿＿＿＿；學校/系級：＿＿＿＿＿＿＿＿')
c = row('提案題目', True); p = simple(c, ''); add_text(p, TITLE, BODY, True)
c = row('主題領域\n(可複選)', True)
p = c.paragraphs[0]; fmt_par(p); add_text(p, '□公共行政 □資訊科技 ■交通運輸 □都市規劃 □觀光行銷')
p = c.add_paragraph(); fmt_par(p); add_text(p, '□環境永續 □多元族群 □產業發展 □教育文化 ■社福醫療')
p = c.add_paragraph(); fmt_par(p); add_text(p, '□其他＿＿＿＿')
c = row('提案摘要'); p = c.paragraphs[0]; fmt_par(p, first=0.99, align=WD_ALIGN_PARAGRAPH.JUSTIFY); add_text(p, ABSTRACT)
c = section('研究動機/\n解決問題'); blocks(c, MOTIVATION)
c = section('研究內容'); blocks(c, CONTENT)
c = section('研究期程')
st_t = c.add_table(rows=1+len(SCHEDULE), cols=2); table_borders(st_t, 4); set_col_widths(st_t, [4.4, LW+CW-0.4-4.4])
for i, h in enumerate(['期間','階段任務']):
    cc = st_t.rows[0].cells[i]; pp = cc.paragraphs[0]; fmt_par(pp, align=WD_ALIGN_PARAGRAPH.CENTER); add_text(pp, h, BODY, True); set_cell_shading(cc,'E7E6E6')
for ri, (a, b) in enumerate(SCHEDULE):
    cant_split(st_t.rows[ri+1])
    pp = st_t.rows[ri+1].cells[0].paragraphs[0]; fmt_par(pp, align=WD_ALIGN_PARAGRAPH.CENTER); add_text(pp, a)
    st_t.rows[ri+1].cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    pp = st_t.rows[ri+1].cells[1].paragraphs[0]; fmt_par(pp, align=WD_ALIGN_PARAGRAPH.JUSTIFY); add_text(pp, b)
p = c.add_paragraph(); fmt_par(p, after=4); add_text(p, SCHEDULE_NOTE, SMALL)
# remove empty first paragraph in schedule cell
fp0 = c.paragraphs[0]
if not fp0.text: fp0._element.getparent().remove(fp0._element)
c = section('研究成果/\n預期效益'); blocks(c, OUTCOME)
c = row('是否使用\nＡＩ工具\n協助創作', True)
p = c.paragraphs[0]; fmt_par(p); add_text(p, '■ 是；工具名稱：＿＿＿＿＿＿＿＿')
p = c.add_paragraph(); fmt_par(p, left=0.99, align=WD_ALIGN_PARAGRAPH.JUSTIFY); add_text(p, '使用方式：' + AI_USE)
p = c.add_paragraph(); fmt_par(p); add_text(p, '□ 否。')
set_col_widths(T, [LW, CW])
# merged cells: fix widths
for rw in T.rows:
    if len(set(id(c._tc) for c in rw.cells))==1: set_cell_width(rw.cells[0], LW+CW)
# cell margins
tblPr = T._tbl.tblPr; mar = OxmlElement('w:tblCellMar')
for side, v in (('left',100),('right',100),('top',40),('bottom',40)):
    e = OxmlElement('w:'+side); e.set(qn('w:w'),str(v)); e.set(qn('w:type'),'dxa'); mar.append(e)
tblPr.append(mar)

# references
p = doc.add_paragraph(); fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6, keep=True); add_text(p, '參考文獻', 16, True)
unused = [k for k in REFS if k not in order]
for k in unused: order.append(k)
for i, k in enumerate(order):
    p = doc.add_paragraph(); fmt_par(p, left=1.0, hanging=1.0, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_text(p, '[%d]\t' % (i+1), 12); add_text(p, REFS[k], 12)
    tabs = p.paragraph_format.tab_stops; tabs.add_tab_stop(Cm(1.0))

out = os.path.join(BASE, '繳交版_研究計畫書.docx')
doc.save(out)
print('saved', out); print('unused refs (appended):', unused)
print('mapping:', {k: i+1 for i, k in enumerate(order)})

# ---- character counts ----
def plain(items):
    s = ''
    for kind, val in items:
        if kind in ('h1','h2','p','li'): s += re.sub(r'\{R\d+\}', '', val)
    return s
def cnt(s): return len(re.sub(r'\s', '', s))
def tabchars(keys):
    n = 0
    for k in keys:
        t = TABLES[k]; n += cnt(t['caption'] + ''.join(t['header']) + ''.join(''.join(r) for r in t['rows']) + re.sub(r'\{R\d+\}','',t['note']))
    return n
def figchars(keys): return sum(cnt(re.sub(r'\{R\d+\}','',FIGS[k][2])) for k in keys)
print('摘要', cnt(ABSTRACT))
print('研究動機', cnt(plain(MOTIVATION)))
print('研究內容 text', cnt(plain(CONTENT)), '+tables', tabchars(list(TABLES)), '+captions', figchars(list(FIGS)))
print('預期效益', cnt(plain(OUTCOME)))
